"""

testcase_ingest.py

Read TestCases_Formatted.xlsx and load it into PostgreSQL (schema_v3.sql):



    one Excel row  ->  1 row in documents

                       1 row in testcase_details   (same doc_id)

                       1..n rows in chunks         (split only if text is too long)



Embeddings are NOT created here: chunks.embedding stays NULL until the
embedding step runs. Chunk limits are counted with the model tokenizer.



Safe to re-run (upsert):

  - unchanged chunk  -> keeps its embedding

  - changed chunk    -> embedding reset to NULL (re-embedded next run)

  - rows/chunks that disappeared from the Excel -> is_active = FALSE

Everything is written in ONE transaction: all or nothing.



Usage (PowerShell):

  pip install openpyxl "psycopg[binary]" tokenizers huggingface_hub

  $env:PG_DSN = "postgresql://postgres:<password>@localhost:5432/ragdb"

  python testcase_ingest.py TestCases_Formatted.xlsx --dry-run     # check first

  python testcase_ingest.py TestCases_Formatted.xlsx               # load

  python testcase_ingest.py TestCases_Formatted.xlsx --only-source-file "ALL MODULES_TCs.xlsx"

"""

import argparse

import hashlib

import json


import os

import re

import sys

from collections import Counter

from pathlib import Path



from openpyxl import load_workbook
from huggingface_hub import hf_hub_download
from tokenizers import Tokenizer



SOURCE_TYPE = "testcase"

AUTHORITY = "style_reference"          # decade-old test cases: format/style only

STRATEGY = "row_per_testcase"

STRATEGY_SPLIT = "row_per_testcase+line_split"



# Count tokens with the Hugging Face tokenizer JSON for the embedding model.
# The embedding prefix is included in the count because it is sent to the model.
DEFAULT_MAX_CHUNK_TOKENS = 400
DEFAULT_TOKENIZER_REPO = "nomic-ai/nomic-embed-text-v2-moe"
DEFAULT_EMBEDDING_PREFIX = "search_document: "
MIN_CHUNK_TOKENS = 32
PART_LINE_RESERVE = 25
MIN_BODY_CHARS = 20


REQUIRED_COLUMNS = ["source_file", "source_sheet", "source_row", "module", "section",

                    "tc_label", "tc_number", "scenario", "preconditions", "steps",

                    "test_data", "expected", "high_level", "tc_source", "owner",

                    "notes", "additional_metadata"]

JUNK_META_KEY = re.compile(r"^(column_\d+|#|S\\.N)$", re.I)   # leftovers from the old sheets





# ---------- small helpers ----------

def sha256(text):

    return hashlib.sha256(text.encode("utf-8")).hexdigest()










def clean(value):

    """Cell -> tidy string: strip, drop tabs/tick emoji, remove blank lines."""

    if value is None:

        return ""

    text = str(value).replace("\r\n", "\n").replace("\t", " ").replace("\u2705", "")

    lines = [re.sub(r"[ ]{2,}", " ", ln).strip() for ln in text.split("\n")]

    return "\n".join(ln for ln in lines if ln).strip()





def or_none(text):

    return text if text else None





def to_bool(text):

    return text.strip().lower() in ("true", "yes", "1")





def to_int(text):

    try:

        return int(float(text))

    except (TypeError, ValueError):

        return None





# ---------- reading the Excel ----------

def read_rows(path):

    wb = load_workbook(path, read_only=True, data_only=True)

    ws = wb["TestCases"] if "TestCases" in wb.sheetnames else wb.worksheets[0]

    it = ws.iter_rows(values_only=True)

    header = [clean(c).lower() for c in next(it)]

    missing = [c for c in REQUIRED_COLUMNS if c not in header]

    if missing:

        sys.exit(f"Excel is missing columns: {missing}")

    idx = {name: header.index(name) for name in REQUIRED_COLUMNS}

    for row in it:

        if not any(c not in (None, "") for c in row):

            continue

        yield {name: clean(row[i]) if i < len(row) else "" for name, i in idx.items()}





# ---------- chunking (one chunk normally; split only when too long) ----------

def shorten(text, limit):

    return text if len(text) <= limit else text[:max(limit - 3, 1)].rstrip() + "..."





def field_lines(label, text, inline):

    parts = text.split("\n")

    if inline:

        return [f"{label}: {parts[0]}"] + parts[1:]

    return [f"{label}:"] + parts





def split_long_line(line, limit):

    """Split one very long line at sentence ends, then at spaces."""

    limit = max(limit, 20)               # never 0 or negative (that made the old version loop forever)

    if len(line) <= limit:

        return [line]

    pieces, cur = [], ""

    for sent in re.split(r"(?<=[.!?;:])\s+", line):

        while len(sent) > limit:

            cut = sent.rfind(" ", 0, limit)

            cut = cut if cut > 0 else limit

            if cur:

                pieces.append(cur)

                cur = ""

            pieces.append(sent[:cut])

            sent = sent[cut:].lstrip()

        if cur and len(cur) + len(sent) + 1 > limit:

            pieces.append(cur)

            cur = sent

        else:

            cur = f"{cur} {sent}".strip()

    if cur:

        pieces.append(cur)

    return pieces





def count_tokens(text, tokenizer, embedding_prefix):
    """Count tokenizer tokens including the prefix and special tokens."""
    return len(tokenizer.encode(embedding_prefix + text, add_special_tokens=True).ids)


def split_line_by_tokens(line, header, max_tokens, tokenizer, embedding_prefix):
    """Split a single long line into pieces that fit with the repeated chunk header."""
    remaining = line
    pieces = []
    # Use a long placeholder so the eventual part label cannot unexpectedly overflow.
    marker = "(part 999999 of 999999)"
    while remaining:
        candidate = f"{header}\n{marker}\n{remaining}"
        if count_tokens(candidate, tokenizer, embedding_prefix) <= max_tokens:
            pieces.append(remaining)
            break

        lo, hi = 1, len(remaining)
        best = 0
        while lo <= hi:
            mid = (lo + hi) // 2
            test = f"{header}\n{marker}\n{remaining[:mid]}"
            if count_tokens(test, tokenizer, embedding_prefix) <= max_tokens:
                best = mid
                lo = mid + 1
            else:
                hi = mid - 1

        if best <= 0:
            raise ValueError(
                "The module/scenario header alone exceeds --max-chunk-tokens. "
                "Increase the token limit or shorten the test-case header."
            )

        # Prefer a word boundary near the end of the largest safe substring.
        cut = best
        if best < len(remaining):
            boundary = remaining.rfind(" ", max(0, int(best * 0.65)), best)
            if boundary > 0:
                cut = boundary
        piece = remaining[:cut].rstrip()
        if not piece:  # avoid an infinite loop on whitespace-only content
            piece = remaining[:best]
            cut = best
        pieces.append(piece)
        remaining = remaining[cut:].lstrip()
    return pieces


def build_chunks(module_line, scenario, other_field_lines, full_content,
                 max_tokens, tokenizer, embedding_prefix):
    """Keep short test cases whole; otherwise split into tokenizer-bounded chunks."""
    if count_tokens(full_content, tokenizer, embedding_prefix) <= max_tokens:
        return [full_content]

    header_module = shorten(module_line, max(60, len(module_line) // 2))
    first_line = scenario.split("\n")[0]
    short_sc = shorten(first_line, min(150, max(40, len(first_line) // 2)))
    header = f"{header_module}\nScenario: {short_sc}"

    body = []
    if scenario != short_sc:
        body += field_lines("Scenario (full)", scenario, inline=False)
    body += other_field_lines

    # Split individual lines first. Every piece is checked with the real tokenizer.
    lines = []
    for line in body:
        if count_tokens(f"{header}\n(part 999999 of 999999)\n{line}", tokenizer, embedding_prefix) <= max_tokens:
            lines.append(line)
        else:
            lines.extend(split_line_by_tokens(line, header, max_tokens, tokenizer, embedding_prefix))

    parts = []
    current = []
    for line in lines:
        candidate_lines = current + [line]
        # Reserve a deliberately long part label while packing, then verify final text below.
        candidate = f"{header}\n(part 999999 of 999999)\n" + "\n".join(candidate_lines)
        if current and count_tokens(candidate, tokenizer, embedding_prefix) > max_tokens:
            parts.append(current)
            current = [line]
        else:
            current = candidate_lines
    if current:
        parts.append(current)

    chunks = []
    total = len(parts)
    for i, part in enumerate(parts, 1):
        text = f"{header}\n(part {i} of {total})\n" + "\n".join(part)
        actual = count_tokens(text, tokenizer, embedding_prefix)
        if actual > max_tokens:
            raise ValueError(
                f"Internal chunking error: chunk {i}/{total} has {actual} tokens, "
                f"limit is {max_tokens}."
            )
        chunks.append(text)
    return chunks


# ---------- one Excel row -> records for the 3 tables ----------

def build_records(row, warnings, max_tokens, tokenizer, embedding_prefix):

    scenario = row["scenario"]

    if not scenario:

        warnings.append(f"{row['tc_label'] or '?'}: skipped (scenario is empty)")

        return None



    doc_id = f"tc:{row['tc_label']}"

    module_line = f"Module: {row['source_sheet']}" + (f" > Section: {row['section']}" if row["section"] else "")



    other = []

    if row["preconditions"]:

        other += field_lines("Preconditions", row["preconditions"], inline=True)

    if row["steps"]:

        other += field_lines("Steps", row["steps"], inline=False)

    if row["test_data"]:

        other += field_lines("Test data", row["test_data"], inline=False)

    if row["expected"]:

        other += field_lines("Expected", row["expected"], inline=True)



    content = "\n".join([module_line] + field_lines("Scenario", scenario, inline=True) + other)

    content_hash = sha256(content)

    high_level = to_bool(row["high_level"])



    # extra info: kept for reference, never embedded

    extra = {}

    if row["additional_metadata"]:

        try:

            extra = {k: v for k, v in json.loads(row["additional_metadata"]).items()

                     if not JUNK_META_KEY.match(str(k))}

        except json.JSONDecodeError:

            warnings.append(f"{row['tc_label']}: additional_metadata is not valid JSON (ignored)")



    document = {

        "doc_id": doc_id,

        "source_type": SOURCE_TYPE,

        "title": scenario.split("\n")[0][:120],

        "content": content,

        "content_hash": content_hash,

        "metadata": {"owner": or_none(row["owner"]), "notes": or_none(row["notes"]),

                     "extra": extra, "has_steps": bool(row["steps"]),

                     "has_expected": bool(row["expected"])},

    }

    details = {

        "doc_id": doc_id,

        "file_name": row["source_file"],

        "sheet_name": row["source_sheet"],

        "module": row["module"],

        "section": or_none(row["section"]),

        "tc_label": row["tc_label"],

        "tc_number": or_none(row["tc_number"]),

        "excel_row": to_int(row["source_row"]),

        "tc_source": or_none(row["tc_source"]),

        "high_level": high_level,

        "authority": AUTHORITY,

        "preconditions": or_none(row["preconditions"]),

        "scenario": scenario,

        "steps": or_none(row["steps"]),

        "test_data": or_none(row["test_data"]),

        "expected": or_none(row["expected"]),

    }



    texts = build_chunks(module_line, scenario, other, content, max_tokens, tokenizer, embedding_prefix)

    chunks = []

    for i, text in enumerate(texts):

        chunks.append({

            "id": f"{doc_id}:c{i}",

            "doc_id": doc_id,

            "source_type": SOURCE_TYPE,

            "chunk_index": i,

            "content": text,

            "content_hash": sha256(text),

            "chunk_strategy": STRATEGY if len(texts) == 1 else STRATEGY_SPLIT,

            "token_count": count_tokens(text, tokenizer, embedding_prefix),

            "metadata": {"part_index": i + 1, "part_count": len(texts),

                         "module": row["module"], "high_level": high_level},

        })

    return document, details, chunks





# ---------- SQL ----------

UPSERT_DOCUMENT = """

INSERT INTO documents (doc_id, source_type, title, content, content_hash, metadata)

VALUES (%(doc_id)s, %(source_type)s, %(title)s, %(content)s, %(content_hash)s, %(metadata)s)

ON CONFLICT (doc_id) DO UPDATE SET

    title = EXCLUDED.title, content = EXCLUDED.content,

    content_hash = EXCLUDED.content_hash, metadata = EXCLUDED.metadata,

    is_active = TRUE, ingested_at = now()

"""



UPSERT_DETAILS = """

INSERT INTO testcase_details (doc_id, file_name, sheet_name, module, section, tc_label,

    tc_number, excel_row, tc_source, high_level, authority, preconditions, scenario,

    steps, test_data, expected)

VALUES (%(doc_id)s, %(file_name)s, %(sheet_name)s, %(module)s, %(section)s, %(tc_label)s,

    %(tc_number)s, %(excel_row)s, %(tc_source)s, %(high_level)s, %(authority)s,

    %(preconditions)s, %(scenario)s, %(steps)s, %(test_data)s, %(expected)s)

ON CONFLICT (doc_id) DO UPDATE SET

    file_name = EXCLUDED.file_name, sheet_name = EXCLUDED.sheet_name,

    module = EXCLUDED.module, section = EXCLUDED.section, tc_label = EXCLUDED.tc_label,

    tc_number = EXCLUDED.tc_number, excel_row = EXCLUDED.excel_row,

    tc_source = EXCLUDED.tc_source, high_level = EXCLUDED.high_level,

    authority = EXCLUDED.authority, preconditions = EXCLUDED.preconditions,

    scenario = EXCLUDED.scenario, steps = EXCLUDED.steps,

    test_data = EXCLUDED.test_data, expected = EXCLUDED.expected

"""



UPSERT_CHUNK = """

INSERT INTO chunks (id, doc_id, source_type, chunk_index, content, content_hash,

                    chunk_strategy, token_count, metadata)

VALUES (%(id)s, %(doc_id)s, %(source_type)s, %(chunk_index)s, %(content)s, %(content_hash)s,

        %(chunk_strategy)s, %(token_count)s, %(metadata)s)

ON CONFLICT (id) DO UPDATE SET

    chunk_index = EXCLUDED.chunk_index, content = EXCLUDED.content,

    chunk_strategy = EXCLUDED.chunk_strategy, token_count = EXCLUDED.token_count,

    metadata = EXCLUDED.metadata, is_active = TRUE,

    embedding       = CASE WHEN chunks.content_hash = EXCLUDED.content_hash

                           THEN chunks.embedding ELSE NULL END,

    embedding_model = CASE WHEN chunks.content_hash = EXCLUDED.content_hash

                           THEN chunks.embedding_model ELSE NULL END,

    content_hash = EXCLUDED.content_hash

"""



# documents that were loaded before from these files but are gone now

DEACTIVATE_DOCS = """

UPDATE documents d SET is_active = FALSE

FROM testcase_details t

WHERE t.doc_id = d.doc_id AND t.file_name = ANY(%s)

  AND NOT (d.doc_id = ANY(%s)) AND d.is_active

"""

# chunks of loaded documents that no longer exist (e.g. 3 parts became 2)

DEACTIVATE_STALE_CHUNKS = """

UPDATE chunks SET is_active = FALSE

WHERE is_active AND doc_id = ANY(%s) AND NOT (id = ANY(%s))

"""

# chunks that belong to deactivated documents

DEACTIVATE_ORPHAN_CHUNKS = """

UPDATE chunks c SET is_active = FALSE

FROM documents d WHERE c.doc_id = d.doc_id AND NOT d.is_active AND c.is_active

"""

SUMMARY = """

SELECT (SELECT count(*) FROM documents WHERE source_type = 'testcase' AND is_active),

       (SELECT count(*) FROM testcase_details t JOIN documents d USING (doc_id) WHERE d.is_active),

       (SELECT count(*) FROM chunks WHERE source_type = 'testcase' AND is_active),

       (SELECT count(*) FROM chunks WHERE source_type = 'testcase' AND embedding IS NULL),

       (SELECT count(*) FROM chunks WHERE source_type = 'testcase' AND NOT is_active)

"""





def batches(items, size):

    for i in range(0, len(items), size):

        yield items[i:i + size]





def write_to_postgres(documents, details, chunks, files, batch_size):

    import psycopg

    from psycopg.types.json import Jsonb



    for d in documents:

        d["metadata"] = Jsonb(d["metadata"])

    for c in chunks:

        c["metadata"] = Jsonb(c["metadata"])



    doc_ids = [d["doc_id"] for d in documents]

    chunk_ids = [c["id"] for c in chunks]



    with psycopg.connect(os.environ["PG_DSN"]) as conn:     # commit on success, rollback on error

        with conn.cursor() as cur:

            # order matters: documents first (the other two tables point to it)

            for b in batches(documents, batch_size):

                cur.executemany(UPSERT_DOCUMENT, b)

            for b in batches(details, batch_size):

                cur.executemany(UPSERT_DETAILS, b)

            for b in batches(chunks, batch_size):

                cur.executemany(UPSERT_CHUNK, b)



            cur.execute(DEACTIVATE_DOCS, (files, doc_ids))

            gone_docs = cur.rowcount

            cur.execute(DEACTIVATE_STALE_CHUNKS, (doc_ids, chunk_ids))

            gone_chunks = cur.rowcount

            cur.execute(DEACTIVATE_ORPHAN_CHUNKS)

            gone_chunks += cur.rowcount



            cur.execute(SUMMARY)

            n_docs, n_details, n_chunks, pending, inactive = cur.fetchone()

    print(f"\nLoaded into PostgreSQL (active testcase rows):")

    print(f"  documents={n_docs}  testcase_details={n_details}  chunks={n_chunks}")

    print(f"  waiting_for_embedding={pending} (active and inactive)   inactive chunks={inactive}")

    if gone_docs or gone_chunks:

        print(f"  deactivated: {gone_docs} documents, {gone_chunks} chunks (no longer in the Excel)")





# ---------- main ----------

def main():

    ap = argparse.ArgumentParser(description="Ingest test cases from Excel into PostgreSQL")

    ap.add_argument("excel", help="path to TestCases_Formatted.xlsx")

    ap.add_argument("--dry-run", action="store_true", help="parse and report only, no database")

    ap.add_argument("--only-source-file", help='load only rows whose source_file equals this, e.g. "ALL MODULES_TCs.xlsx"')

    ap.add_argument("--batch-size", type=int, default=500)

    ap.add_argument("--max-chunk-tokens", type=int, default=DEFAULT_MAX_CHUNK_TOKENS,
                    help=f"maximum tokenizer-counted tokens per chunk, default {DEFAULT_MAX_CHUNK_TOKENS}")
    ap.add_argument("--tokenizer-repo", default=DEFAULT_TOKENIZER_REPO,
                    help=f"Hugging Face repository containing tokenizer.json (default: {DEFAULT_TOKENIZER_REPO})")
    ap.add_argument("--embedding-prefix", default=DEFAULT_EMBEDDING_PREFIX,
                    help=f"prefix sent to the embedding model (default: {DEFAULT_EMBEDDING_PREFIX!r})")
    # Compatibility only: old commands can still pass this, but it is ignored.
    ap.add_argument("--chars-per-token", type=float, default=None, help=argparse.SUPPRESS)
    args = ap.parse_args()
    if args.max_chunk_tokens < MIN_CHUNK_TOKENS:
        sys.exit(f"--max-chunk-tokens must be at least {MIN_CHUNK_TOKENS}")
    if args.max_chunk_tokens > 480:
        sys.exit("--max-chunk-tokens must be 480 or less to leave a safety margin below the model's 512-token context window.")

    if not Path(args.excel).exists():

        sys.exit(f"File not found: {args.excel}")

    if not args.dry_run and "PG_DSN" not in os.environ:

        sys.exit("Set PG_DSN first, e.g.  $env:PG_DSN = 'postgresql://postgres:<password>@localhost:5432/ragdb'")



    try:
        tokenizer_path = hf_hub_download(repo_id=args.tokenizer_repo, filename="tokenizer.json")
        tokenizer = Tokenizer.from_file(tokenizer_path)
    except Exception as exc:
        sys.exit(
            "Could not load tokenizer.json from Hugging Face. Check internet access and "
            "install dependencies with: pip install tokenizers huggingface_hub\n"
            f"Details: {exc}"
        )

    warnings, documents, details, chunks = [], [], [], []

    seen = Counter()

    for row in read_rows(args.excel):

        if args.only_source_file and row["source_file"] != args.only_source_file:

            continue

        if not row["tc_label"]:

            warnings.append(f"{row['source_sheet']} row {row['source_row']}: skipped (no tc_label)")

            continue

        seen[row["tc_label"]] += 1

        rec = build_records(row, warnings, args.max_chunk_tokens, tokenizer, args.embedding_prefix)

        if rec:

            documents.append(rec[0]); details.append(rec[1]); chunks.extend(rec[2])



    dupes = [k for k, v in seen.items() if v > 1]

    if dupes:

        sys.exit(f"tc_label must be unique, but these repeat: {dupes[:5]} ... ({len(dupes)} total)")

    if not documents:

        sys.exit("Nothing to load.")



    files = sorted({d["file_name"] for d in details})

    parts = Counter(c["doc_id"] for c in chunks)

    split = sum(1 for v in parts.values() if v > 1)

    biggest_chars = max(len(c["content"]) for c in chunks)
    biggest_tokens = max(c["token_count"] for c in chunks)
    over = sum(1 for c in chunks if c["token_count"] > args.max_chunk_tokens)
    if over:
        sys.exit(f"BUG: {over} chunks exceed the limit of {args.max_chunk_tokens} tokens.")
    by_file = Counter(d["file_name"] for d in details)



    print(f"Read {len(documents)} test cases -> {len(documents)} documents, "

          f"{len(details)} testcase_details, {len(chunks)} chunks")

    print(f"  split into several chunks: {split} | longest chunk: {biggest_tokens} tokens "
          f"(limit {args.max_chunk_tokens}; longest text {biggest_chars} characters)")
    print(f"  tokenizer: {args.tokenizer_repo} | prefix counted: {args.embedding_prefix!r}")
    print(f"  chunks per test case: {dict(sorted(Counter(parts.values()).items()))}")

    for f, n in by_file.items():

        print(f"  {f}: {n}")

    for w in warnings[:10]:

        print("  WARN", w)

    if len(warnings) > 10:

        print(f"  ... {len(warnings) - 10} more warnings")



    if args.dry_run:

        print("\n--dry-run: nothing written. First document:")

        print(json.dumps({k: documents[0][k] for k in ("doc_id", "title", "content_hash")}, indent=2))

        return

    write_to_postgres(documents, details, chunks, files, args.batch_size)





if __name__ == "__main__":

    main()